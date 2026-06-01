"""Intercoder – report parsing (docx / xlsx / html → impulse DataFrame)."""
from pathlib import Path

import pandas as pd


def parse_report_impulses(file_path):
    """Dispatch on file extension. Supports .docx, .xlsx, .html/.htm."""
    ext = Path(file_path).suffix.lower()
    if ext == ".docx":
        return _parse_report_impulses_docx(file_path)
    if ext == ".xlsx":
        return _parse_report_impulses_xlsx(file_path)
    if ext in (".html", ".htm"):
        return _parse_report_impulses_html(file_path)
    raise ValueError("unsupported_format")


def _rows_to_impulses_df(matrix):
    rows = []
    for cells in matrix:
        ncols = len(cells)
        if ncols < 3:
            continue
        if ncols >= 4:
            sprecher = cells[1]
            impuls = cells[-2]
            code = cells[-1]
        else:
            sprecher = ""
            impuls = cells[-2]
            code = cells[-1]
        if not impuls or not str(impuls).strip():
            continue
        rows.append({"Sprecher": str(sprecher).strip(),
                     "Impuls": str(impuls).strip(),
                     "Shortcode": str(code).strip()})
    return pd.DataFrame(rows, columns=["Sprecher", "Impuls", "Shortcode"])


def _parse_report_impulses_docx(docx_file_path):
    from docx import Document

    doc = Document(docx_file_path)
    target = None
    for tbl in doc.tables:
        try:
            first_header = tbl.cell(0, 0).text.strip()
        except Exception:
            continue
        if first_header == "#":
            target = tbl
            break
    if target is None:
        raise ValueError("no_impulse_table")

    if len(target.columns) < 3:
        raise ValueError("no_impulse_table")

    matrix = [[c.text.strip() for c in row.cells]
              for ri, row in enumerate(target.rows) if ri > 0]
    df = _rows_to_impulses_df(matrix)
    if df.empty:
        raise ValueError("no_impulse_table")
    return df


def _parse_report_impulses_xlsx(xlsx_file_path):
    try:
        import openpyxl  # noqa: F401
    except ImportError as e:
        raise ValueError("xlsx_unavailable") from e
    xl = pd.ExcelFile(xlsx_file_path)
    target_df = None
    for name in xl.sheet_names:
        df = pd.read_excel(xlsx_file_path, sheet_name=name)
        if df.shape[1] < 3:
            continue
        if str(df.columns[0]).strip() == "#":
            target_df = df
            break
    if target_df is None:
        raise ValueError("no_impulse_table")
    matrix = [["" if pd.isna(c) else c for c in row.tolist()]
              for _, row in target_df.iterrows()]
    df = _rows_to_impulses_df(matrix)
    if df.empty:
        raise ValueError("no_impulse_table")
    return df


def _parse_report_impulses_html(html_file_path):
    try:
        tables = pd.read_html(html_file_path)
    except Exception as e:
        raise ValueError("no_impulse_table") from e
    target = None
    for df in tables:
        if df.shape[1] < 3:
            continue
        if str(df.columns[0]).strip() == "#":
            target = df
            break
    if target is None:
        raise ValueError("no_impulse_table")
    matrix = [["" if pd.isna(c) else c for c in row.tolist()]
              for _, row in target.iterrows()]
    df = _rows_to_impulses_df(matrix)
    if df.empty:
        raise ValueError("no_impulse_table")
    return df
